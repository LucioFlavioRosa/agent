import logging
import io
import pytest
from backend.app.utils.docx_parser import extrair_texto_docx_em_memoria
from unittest.mock import patch

class DummyDocx:
    """Classe dummy para simular bytes de um .docx válido."""
    @staticmethod
    def get_bytes():
        # Gera um arquivo docx simples em memória
        from docx import Document
        doc = Document()
        doc.add_paragraph("Primeiro parágrafo de teste.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0,0).text = "Celula 1"
        table.cell(0,1).text = "Celula 2"
        table.cell(1,0).text = "Celula 3"
        table.cell(1,1).text = "Celula 4"
        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()

@pytest.fixture
def log_capture():
    logger = logging.getLogger("mcp_docx_parser")
    old_level = logger.level
    logger.setLevel(logging.DEBUG)
    log_handler = logging.StreamHandler()
    log_handler.setLevel(logging.DEBUG)
    log_records = []
    def emit(record):
        log_records.append(record)
    log_handler.emit = emit
    logger.addHandler(log_handler)
    yield log_records
    logger.removeHandler(log_handler)
    logger.setLevel(old_level)

@pytest.mark.asyncio
async def test_extrair_texto_docx_em_memoria_logs_success(log_capture):
    file_bytes = DummyDocx.get_bytes()
    job_id = "job123"
    company_id = "comp456"
    project_id = "proj789"
    resultado = extrair_texto_docx_em_memoria(file_bytes, job_id=job_id, company_id=company_id, project_id=project_id)
    # Verifica logs essenciais
    logs_text = [record.getMessage() for record in log_capture]
    assert any("Entrada no extrair_texto_docx_em_memoria" in msg for msg in logs_text)
    assert any("Iniciando iteração de blocos" in msg for msg in logs_text)
    assert any("Processando parágrafo" in msg for msg in logs_text)
    assert any("Início da tabela" in msg for msg in logs_text)
    assert any("Fim da tabela" in msg for msg in logs_text)
    assert any("Extração finalizada" in msg for msg in logs_text)
    # Verifica contexto
    for campo in [job_id, company_id, project_id]:
        assert any(campo in msg for msg in logs_text)

@pytest.mark.asyncio
async def test_extrair_texto_docx_em_memoria_logs_error(log_capture):
    # Simula erro passando bytes inválidos
    file_bytes = b"not_a_valid_docx"
    job_id = "job_error"
    company_id = "comp_error"
    project_id = "proj_error"
    with pytest.raises(Exception):
        extrair_texto_docx_em_memoria(file_bytes, job_id=job_id, company_id=company_id, project_id=project_id)
    logs_text = [record.getMessage() for record in log_capture]
    assert any("ERRO ao extrair texto do DOCX" in msg for msg in logs_text)
    # Verifica stack trace foi logado
    error_records = [record for record in log_capture if record.levelname == "ERROR"]
    assert error_records, "Deve haver pelo menos um log de erro"
    # Verifica contexto
    for campo in [job_id, company_id, project_id]:
        assert any(campo in msg for msg in logs_text)
