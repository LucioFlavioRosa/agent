import io
import pytest
from unittest import mock
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

# Supondo que extract_text_from_docx esteja em backend/services/docx_parser.py
from backend.services.docx_parser import extract_text_from_docx

def create_docx_with_text(text):
    doc = Document()
    doc.add_paragraph(text)
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def create_empty_docx():
    doc = Document()
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def test_extract_text_from_valid_docx():
    file_stream = create_docx_with_text("Texto de teste para extração.")
    result = extract_text_from_docx(file_stream)
    assert "Texto de teste para extração." in result

def test_extract_text_from_empty_docx():
    file_stream = create_empty_docx()
    result = extract_text_from_docx(file_stream)
    assert result.strip() == ""

def test_extract_text_from_invalid_docx():
    # Cria um arquivo inválido (não é um docx de verdade)
    invalid_stream = io.BytesIO(b"not a real docx content")
    with pytest.raises((PackageNotFoundError, Exception)):
        extract_text_from_docx(invalid_stream)
